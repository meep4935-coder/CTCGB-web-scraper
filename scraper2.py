""" Program Web Scraper for CTCGB
    To run the program, click run, then in console type; uv run python directory.
    Designed to determine whether clean tech programs are open or closed.
    Created 2026-07-09 author @ kwong s"""
import os
import re
import time
import sys
import random
import subprocess
import pandas as pd
from bs4 import BeautifulSoup

# Tkinter imports for GUI
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# --- CONFIGURATION & KEYWORDS ---

PRIORITY_ELEMENTS = [
    'h1', 'h2', 'h3', 'strong', 'b',
    '.alert', '.notice', '.banner', '.status', '.message', '.warning',
    '.mw-brand', '.callout', '.hero', '.closed', '.archived'
]

CLOSED_PATTERNS = [
    r'\bclosed\b',
    r'\bpassed\s+deadline\b',
    r'\bdeadline\s+has\s+passed\b',
    r'\bpaused\b',
    r'\bno\s+longer\s+accepting\b',
    r'\bnot\s+accepting\b',
    r'\bintake\s+is\s+closed\b',
    r'\bclosed\s+for\s+intake\b',
    r'\bproposals\s+(?:is|are)\s+now\s+closed\b',
    r'\bended\b',
    r'\bfully\s+allocated\b',
]

# Zone-restricted closed patterns: these words ONLY count as "Closed" if found
# within the specified HTML selectors. Finding them in body text alone is ignored.
ZONE_RESTRICTED_CLOSED_PATTERNS = {
    r'\barchived\b': ['h2'],
    r'\bsunset\b':   ['h1', '.alert'],
}

OPEN_PATTERNS = [
    r'\bopen\s+for\s+intake\b',
    r'\baccepting\s+applications\b',
    r'\bapply\s+now\b',
    r'\bsubmit\s+your\s+application\b',
    r'\bopen\s+until\b',
    r'\bapply\s+today\b',
    r'\bintake\s+is\s+open\b',
    r'\bnow\s+open\b',
    r'\baccepting\s+proposals\b',
    r'\bopen\s+call\b',
]

INQUIRY_PATTERNS = [
    r'\bcontact\s+(?:us\s+)?to\s+apply\b',
    r'\bbook\s+a\s+(?:meeting|call|consultation|demo)\b',
    r'\bschedule\s+a\s+(?:meeting|call|consultation|demo)\b',
    r'\breach\s+out\s+to\b',
    r'\bget\s+in\s+touch\s+(?:with|to)\b',
    r'\bcontact\s+(?:our\s+team|us\s+for\s+more\s+info)\b',
    r'\bemail\s+us\s+(?:at|to)\b',
    r'\bcall\s+us\s+(?:at|to)\b',
    r'\blet\'s\s+talk\b',
    r'\bspeak\s+(?:to|with)\s+(?:us|our\s+team|an\s+expert)\b',
    r'\binquire\s+(?:about|here)\b',
]

EMAIL_PATTERN = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
PHONE_PATTERN = r'\b(?:\+?1[-. ]?)?\(?([0-9]{3})\)?[-. ]?([0-9]{3})[-. ]?([0-9]{4})\b'

FINANCIAL_PATTERNS = [
    r'\btax\s+credit\b',
    r'\btax\b',
    r'\bfinancing\b',
    r'\bbanking\b',
    r'\bloan\b',
    r'\bgrant\b',
    r'\binvestment\s+tax\s+credit\b',
    r'\bitc\b',
    r'\binvestments\b',
    r'\bexpenses\b',
    r'\bcapital\b',
    r'\ballowence\b',
]

USER_AGENTS = [
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/16.1 Safari/605.1.15',
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:109.0) Gecko/20100101 Firefox/115.0',
    'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/114.0.0.0 Safari/537.36',
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36 Edg/120.0.0.0'
]


def fetch_page(url, browser, timeout=30):
    """Fetches the HTML content of the page using playwright."""
    current_ua = random.choice(USER_AGENTS)
    
    # Create a new context for each request to rotate user agent and isolate cookies
    context = browser.new_context(
        user_agent=current_ua,
        viewport={'width': 1920, 'height': 1080},
    )
    page = context.new_page()
    try:
        # Wait until networkidle for maximum accuracy
        response = page.goto(url, timeout=timeout * 1000, wait_until="networkidle")
        
        if response is None:
            return None, "Failed to load webpage (No response)"
        
        if response.status == 404:
            return None, "404 Not Found"
            
        if response.status == 403:
            return None, "403 Forbidden"
            
        html = page.content()
        return html, None
    except Exception as e:
        err_str = str(e)
        if "Timeout" in err_str:
            return None, f"Timeout after {timeout}s"
        return None, err_str
    finally:
        context.close()


def classify_page(html, url):
    """Determines if a page is Open, Closed, or Manual Review based on hierarchical heuristics."""
    soup = BeautifulSoup(html, 'html.parser')
    
    # 1. Clean up script and style elements
    for script in soup(["script", "style"]):
        script.decompose()

    # 2. Extract text from high-priority elements
    priority_texts = []
    for selector in PRIORITY_ELEMENTS:
        try:
            elements = soup.select(selector)
            for el in elements:
                text = el.get_text(separator=' ', strip=True)
                if text:
                    priority_texts.append((selector, text))
        except Exception:
            pass

    # Standard body text
    body_text = soup.get_text(separator=' ', strip=True)

    # --- STEP 0: CHECK FOR FINANCIAL PROGRAM INDICATORS (overrides Closed) ---
    # Only check the URL and page <title> — NOT the full body text,
    # to avoid false positives on generic financial/banking websites.
    financial_flag = False
    financial_clue = ""
    page_title_tag = soup.find('title')
    page_title = page_title_tag.get_text(strip=True) if page_title_tag else ""
    financial_search_text = url + " " + page_title
    for pattern in FINANCIAL_PATTERNS:
        match = re.search(pattern, financial_search_text, re.IGNORECASE)
        if match:
            financial_flag = True
            financial_clue = f"Financial program indicator '{match.group(0)}' found in URL/title"
            break

    # --- STEP 1: CHECK FOR INQUIRY/ACTION INDICATORS ---
    for pattern in INQUIRY_PATTERNS:
        match = re.search(pattern, body_text, re.IGNORECASE)
        if match:
            start = max(0, match.start() - 50)
            end = min(len(body_text), match.end() + 50)
            snippet = body_text[start:end].replace('\n', ' ').strip()
            return "Manual Review", f"Actionable inquiry found '{match.group(0)}': '...{snippet}...'"

    # --- STEP 2: CHECK FOR OPEN & CLOSED INDICATORS ---
    found_closed = False
    closed_clue = ""

    # 2a. Zone-restricted patterns (archived → h2 only; sunset → h1/.alert only)
    for pattern, allowed_zones in ZONE_RESTRICTED_CLOSED_PATTERNS.items():
        for zone in allowed_zones:
            try:
                for el in soup.select(zone):
                    zone_text = el.get_text(separator=' ', strip=True)
                    match = re.search(pattern, zone_text, re.IGNORECASE)
                    if match:
                        found_closed = True
                        closed_clue = f"Found zone-restricted closed indicator '{match.group(0)}' in required zone ({zone}): '{zone_text[:100]}...'"
                        break
            except Exception:
                pass
            if found_closed:
                break
        if found_closed:
            break

    # 2b. Standard closed patterns — priority zones first, then full body
    if not found_closed:
        for tag, text in priority_texts:
            for pattern in CLOSED_PATTERNS:
                match = re.search(pattern, text, re.IGNORECASE)
                if match:
                    found_closed = True
                    closed_clue = f"Found closed indicator '{match.group(0)}' in priority zone ({tag}): '{text[:100]}...'"
                    break
            if found_closed:
                break

    if not found_closed:
        for pattern in CLOSED_PATTERNS:
            match = re.search(pattern, body_text, re.IGNORECASE)
            if match:
                found_closed = True
                start = max(0, match.start() - 50)
                end = min(len(body_text), match.end() + 50)
                snippet = body_text[start:end].replace('\n', ' ').strip()
                closed_clue = f"Found closed indicator '{match.group(0)}' in body text: '...{snippet}...'"
                break

    # 2c. Open patterns — priority zones first, then full body
    found_open = False
    open_clue = ""
    for tag, text in priority_texts:
        for pattern in OPEN_PATTERNS:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                found_open = True
                open_clue = f"Found open indicator '{match.group(0)}' in priority zone ({tag}): '{text[:100]}...'"
                break
        if found_open:
            break

    if not found_open:
        for pattern in OPEN_PATTERNS:
            match = re.search(pattern, body_text, re.IGNORECASE)
            if match:
                found_open = True
                start = max(0, match.start() - 50)
                end = min(len(body_text), match.end() + 50)
                snippet = body_text[start:end].replace('\n', ' ').strip()
                open_clue = f"Found open indicator '{match.group(0)}' in body text: '...{snippet}...'"
                break

    # --- Resolve final status ---
    # Financial programs override Closed — they stay as Manual Review unless also explicitly Open
    if financial_flag:
        if found_open:
            return "Open", open_clue
        return "Manual Review", f"{financial_clue} (no explicit open indicator)"

    if found_open and found_closed:
        return "Manual Review", f"Conflicting signals (both Open and Closed found). Closed clue: {closed_clue} | Open clue: {open_clue}"
    elif found_closed:
        return "Closed", closed_clue
    elif found_open:
        return "Open", open_clue

    # --- STEP 3: FALLBACK TO BASIC CONTACT INFO ---
    emails = re.findall(EMAIL_PATTERN, body_text)
    phones = re.findall(PHONE_PATTERN, body_text)
    
    contact_in_url = 'contact' in url.lower()
    contact_links = soup.find_all('a', href=lambda href: href and 'contact' in href.lower())

    if emails or phones or contact_in_url or contact_links:
        clue = []
        if emails:
            clue.append(f"Emails found: {list(set(emails))[:2]}")
        if phones:
            formatted_phones = [f"({p[0]})-{p[1]}-{p[2]}" for p in phones]
            clue.append(f"Phones found: {list(set(formatted_phones))[:2]}")
        if contact_links or contact_in_url:
            clue.append("Contact links/URL found")
        return "Manual Review", f"No explicit open/closed status; Contact details found: {'; '.join(clue)}"

    return "Manual Review", "No explicit status keywords or contact details found on page."


def prompt_for_column(columns):
    """Displays a graphical window to select the URL column from a dropdown menu."""
    selected_column = [None]
    
    root = tk.Tk()
    root.title("Select URL Column")
    root.geometry("400x150")
    root.resizable(False, False)
    
    root.lift()
    root.attributes('-topmost', True)
    
    root.update_idletasks()
    x = (root.winfo_screenwidth() // 2) - 200
    y = (root.winfo_screenheight() // 2) - 75
    root.geometry(f"+{x}+{y}")

    label = ttk.Label(root, text="Select the column containing website URLs:", font=("Arial", 11))
    label.pack(pady=15)

    default_idx = 0
    guesses = ['program', 'url', 'link', 'website', 'address', 'page']
    for idx, col in enumerate(columns):
        if any(g in str(col).lower() for g in guesses):
            default_idx = idx
            break

    # Cast column names to strings to prevent Combobox rendering errors
    combobox = ttk.Combobox(root, values=[str(c) for c in columns], state="readonly", width=35)
    combobox.pack(pady=5)
    if len(columns) > 0:
        combobox.current(default_idx)

    def on_submit():
        val = combobox.get()
        # Map string back to the original column type to prevent KeyError in pandas later
        for col in columns:
            if str(col) == val:
                selected_column[0] = col
                break
        root.destroy()

    button = ttk.Button(root, text="Confirm", command=on_submit)
    button.pack(pady=15)

    root.mainloop()
    return selected_column[0]


def main():
    # 1. Prompt for Input Excel File
    temp_root = tk.Tk()
    temp_root.withdraw()
    temp_root.lift()
    temp_root.attributes('-topmost', True)
    
    input_file = filedialog.askopenfilename(
        parent=temp_root,
        title="Select Input Excel File",
        filetypes=[("Excel Files", "*.xlsx *.xls")]
    )
    temp_root.destroy()
    
    if not input_file:
        print("No input file selected. Exiting.")
        sys.exit(0)

    try:
        df = pd.read_excel(input_file)
    except Exception as e:
        temp_root = tk.Tk()
        temp_root.withdraw()
        messagebox.showerror("Error", f"Could not read Excel file:\n{e}")
        temp_root.destroy()
        sys.exit(1)

    url_column = prompt_for_column(df.columns)
    if url_column is None:
        print("No URL column selected. Exiting.")
        sys.exit(0)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(input_file, data_only=False)
        sheet = wb.active
        
        col_idx = None
        for c in range(1, sheet.max_column + 1):
            if str(sheet.cell(row=1, column=c).value) == str(url_column):
                col_idx = c
                break
        
        if col_idx is not None:
            extracted_links = 0
            for r in range(2, sheet.max_row + 1):
                cell = sheet.cell(row=r, column=col_idx)
                df_idx = r - 2
                if df_idx < len(df):
                    if cell.hyperlink and cell.hyperlink.target:
                        df.at[df_idx, url_column] = cell.hyperlink.target
                        extracted_links += 1
            print(f"Extracted {extracted_links} underlying cell hyperlinks using openpyxl.")
    except Exception as e:
        print(f"Warning: Could not extract cell hyperlinks via openpyxl ({e}). Using cell text values.")

    base_dir, orig_name = os.path.split(input_file)
    name_part, ext_part = os.path.splitext(orig_name)
    suggested_output_name = f"{name_part}_classified{ext_part}"
    
    temp_root = tk.Tk()
    temp_root.withdraw()
    temp_root.lift()
    temp_root.attributes('-topmost', True)
    
    output_file = filedialog.asksaveasfilename(
        parent=temp_root,
        title="Save Output Excel File As",
        initialdir=base_dir,
        initialfile=suggested_output_name,
        filetypes=[("Excel Files", "*.xlsx")]
    )
    temp_root.destroy()
    
    if not output_file:
        print("No output file selected. Exiting.")
        sys.exit(0)

    df_scrape = df[df[url_column].astype(str).str.startswith(('http://', 'https://'))]
    total_urls = len(df_scrape)
    
    if total_urls == 0:
        temp_root = tk.Tk()
        temp_root.withdraw()
        messagebox.showwarning("Warning", f"No valid URLs starting with http:// or https:// found in column '{url_column}'.")
        temp_root.destroy()
        sys.exit(0)

    # Playwright automatic setup
    try:
        # type: ignore
        from playwright.sync_api import sync_playwright
    except ImportError:
        err_root = tk.Tk()
        err_root.withdraw()
        messagebox.showerror("Error", "playwright is not installed. Ensure it's in the PEP 723 dependencies block.")
        err_root.destroy()
        sys.exit(1)
        
    print("Ensuring Playwright browsers are installed...")
    try:
        subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except Exception:
        print("Note: Could not automatically run 'playwright install'. You may need to run 'uv run playwright install' manually.")

    progress_window = tk.Tk()
    progress_window.title("Scraping Websites with Browser...")
    progress_window.geometry("450x150")
    progress_window.resizable(False, False)
    progress_window.lift()
    progress_window.attributes('-topmost', True)
    
    progress_window.update_idletasks()
    x = (progress_window.winfo_screenwidth() // 2) - 225
    y = (progress_window.winfo_screenheight() // 2) - 75
    progress_window.geometry(f"+{x}+{y}")

    lbl_status = ttk.Label(progress_window, text=f"Starting scan (0 of {total_urls})...", font=("Arial", 10))
    lbl_status.pack(pady=15)

    progress_bar = ttk.Progressbar(progress_window, orient="horizontal", length=350, mode="determinate")
    progress_bar.pack(pady=5)
    progress_bar["maximum"] = total_urls

    lbl_current_url = ttk.Label(progress_window, text="", font=("Arial", 8), foreground="gray")
    lbl_current_url.pack(pady=10)

    progress_window.update()

    scraped_results = {}

    print("Launching headless Chromium browser...")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        
        for i, (index, row) in enumerate(df_scrape.iterrows()):
            url = str(row[url_column]).strip()
            
            if progress_window.winfo_exists():
                lbl_status.config(text=f"Scraping program {i+1} of {total_urls}...")
                lbl_current_url.config(text=url[:60] + "..." if len(url) > 60 else url)
                progress_bar["value"] = i + 1
                progress_window.update()
            
            print(f"[{i+1}/{total_urls}] Scraping: {url}")
            
            html, error = fetch_page(url, browser)
            
            if error:
                if "404 Not Found" in error:
                    print(f"  -> STATUS: Closed")
                    print(f"  -> CLUE: Page not found (404)")
                    scraped_results[url] = ("Closed", "Page not found (404)", "")
                else:
                    print(f"  -> STATUS: Manual Review")
                    print(f"  -> CLUE: Failed to load webpage ({error})")
                    scraped_results[url] = ("Manual Review", "Failed to load webpage", error)
            else:
                status, clue = classify_page(html, url)
                print(f"  -> STATUS: {status}")
                print(f"  -> CLUE: {clue}")
                scraped_results[url] = (status, clue, "")

            print("-" * 60)
            delay = random.uniform(2.0, 5.0)
            time.sleep(delay)

        browser.close()

    if progress_window.winfo_exists():
        progress_window.destroy()

    import shutil
    try:
        shutil.copy(input_file, output_file)
        
        import openpyxl
        wb = openpyxl.load_workbook(output_file, data_only=False)
        sheet = wb.active
        
        col_url_idx = None
        col_status_idx = None
        col_comments_idx = None
        
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=c).value
            if str(val) == str(url_column):
                col_url_idx = c
            elif val == 'Status':
                col_status_idx = c
            elif val == 'Comments for requested edits':
                col_comments_idx = c
                
        if col_url_idx is None:
            raise ValueError(f"Could not find URL column '{url_column}' in the output sheet.")
            
        if col_status_idx is None:
            col_status_idx = sheet.max_column + 1
            sheet.cell(row=1, column=col_status_idx, value='Status')
        if col_comments_idx is None:
            col_comments_idx = sheet.max_column + 1
            sheet.cell(row=1, column=col_comments_idx, value='Comments for requested edits')
            
        for r in range(2, sheet.max_row + 1):
            cell_url = sheet.cell(row=r, column=col_url_idx)
            url = None
            if cell_url.hyperlink and cell_url.hyperlink.target:
                url = cell_url.hyperlink.target
            elif cell_url.value:
                url = str(cell_url.value).strip()
                
            if url and url in scraped_results:
                status, clue, error = scraped_results[url]
                sheet.cell(row=r, column=col_status_idx, value=status)
                if error:
                    sheet.cell(row=r, column=col_comments_idx, value=f"Scrape Error: {error}")
                else:
                    sheet.cell(row=r, column=col_comments_idx, value=clue)
                    
        wb.save(output_file)
        
        temp_root = tk.Tk()
        temp_root.withdraw()
        temp_root.lift()
        temp_root.attributes('-topmost', True)
        messagebox.showinfo("Success", f"Scraping completed successfully using Playwright!\n\nResults saved to:\n{output_file}")
        temp_root.destroy()
    except Exception as e:
        temp_root = tk.Tk()
        temp_root.withdraw()
        temp_root.lift()
        temp_root.attributes('-topmost', True)
        messagebox.showerror("Error", f"Error writing Excel output:\n{e}")
        temp_root.destroy()


if __name__ == "__main__":
    main()